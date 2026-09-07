"""Structure extraction: headings, section paths, and section-aware chunking.

The property under test throughout is the one STAIR (arXiv:2609.03874) argues
for: a chunk is a whole number of the document's own sections, and it says
which ones. A chunk that silently spans half of one heading and half of the
next is the failure these tests exist to catch.
"""

import io

import pytest

from app.documents import chunk_blocks, extract
from app.structure import (
    PATH_SEPARATOR,
    build_toc,
    detect,
    detect_html,
    detect_markdown,
    detect_plain,
    merge_section_paths,
    section_paths,
)


def paths_of(chunks):
    return [chunk["section_path"] for chunk in chunks]


def test_markdown_headings_build_a_nested_path():
    text = "# Handbook\n\n## Leave\n\nintro\n\n### Sick leave\n\nup to eight days\n"
    headings = detect_markdown(text)
    assert [(h.level, h.title) for h in headings] == [
        (1, "Handbook"), (2, "Leave"), (3, "Sick leave")
    ]
    spans = section_paths(text, headings)
    deepest = spans[-1][2]
    assert deepest == "Handbook > Leave > Sick leave"


def test_a_deeper_heading_does_not_close_its_parent():
    text = "# A\n\n## B\n\n### C\n\ntext\n\n## D\n\nmore\n"
    spans = section_paths(text, detect_markdown(text))
    assert "A > B > C" in [path for _, _, path in spans]
    # D is a sibling of B, so C must not survive into it.
    assert "A > D" in [path for _, _, path in spans]


def test_text_before_the_first_heading_keeps_the_base_path():
    text = "preamble\n\n# Title\n\nbody\n"
    spans = section_paths(text, detect_markdown(text))
    assert spans[0][2] == ""
    assert text[spans[0][0]:spans[0][1]].strip() == "preamble"


def test_fenced_code_is_not_mistaken_for_headings():
    text = "# Real\n\n```\n# not a heading\n## also not\n```\n\nbody\n"
    assert [h.title for h in detect_markdown(text)] == ["Real"]


def test_persian_legal_numbering_is_detected_with_its_nesting():
    text = "فصل اول — کلیات\nمتن\nماده ۳ — معامله کوچک\nمتن\nتبصره ۱\nسه فروشنده\n"
    headings = detect_plain(text)
    assert [(h.level, h.title) for h in headings] == [
        (1, "فصل اول — کلیات"), (2, "ماده ۳ — معامله کوچک"), (3, "تبصره ۱")
    ]
    assert section_paths(text, headings)[-1][2] == "فصل اول — کلیات > ماده ۳ — معامله کوچک > تبصره ۱"


@pytest.mark.parametrize("line", [
    "ماده 12 — تعاریف",
    "Chapter 4",
    "Section 2. Forecast accuracy",
    "4.2.1 Rollback procedure",
])
def test_numbering_conventions_across_scripts(line):
    assert detect_plain(line + "\nbody\n"), f"{line!r} should read as a heading"


@pytest.mark.parametrize("line", [
    "مبلغ کل قرارداد ۲۴۰ میلیون ریال است و در چهار قسط پرداخت می‌شود",
    "In 2026 the network handled 4.18 million shipments across the region",
    "1 shipment was delayed,",
    # A sentence that opens with a heading keyword is still a sentence.
    "بند اول این است که همه کارکنان باید در دفتر حضور داشته باشند",
    "Section 3 of the report was written by the finance team last quarter",
    "3.5 percent of the fleet was replaced during the reporting year",
    "این آیین‌نامه ترتیب انجام معاملات شرکت را تعیین می‌کند.",
])
def test_prose_is_not_mistaken_for_a_heading(line):
    assert not detect_plain(line + "\n"), f"{line!r} is prose, not a heading"


def test_html_headings_survive_instead_of_being_flattened():
    html = "<h1>Release notes</h1><p>intro</p><h2>7.4.0</h2><p>scheduler change</p>"
    blocks = detect_html(html)
    assert any(path == "Release notes" and "intro" in text for text, _, path in blocks)
    assert any(path == "Release notes > 7.4.0" for _, _, path in blocks)


def test_a_heading_stays_in_the_text_of_the_section_it_opens():
    """Otherwise a term that appears only in a heading is lexically unfindable.

    Markdown and plain text keep it because their sections are sliced out of
    the raw string; HTML and Word have to put it back deliberately.
    """
    blocks = detect_html("<h1>Notes</h1><h2>Rollback procedure</h2><p>Run the revert script.</p>")
    body = next(text for text, _, path in blocks if "Rollback" in path)
    assert "Rollback procedure" in body


def test_word_headings_stay_in_the_text_too():
    import io as byte_io

    from docx import Document

    document = Document()
    document.add_heading("Rollback procedure", level=1)
    document.add_paragraph("Run the revert script.")
    buffer = byte_io.BytesIO()
    document.save(buffer)

    chunk = chunk_blocks(extract("runbook.docx", buffer.getvalue()), 900, 100)[0]
    assert "Rollback procedure" in chunk["content"]
    assert "Rollback procedure" in chunk["section_path"]


def test_html_without_headings_still_yields_its_text():
    blocks = detect_html("<p>just a paragraph</p>")
    assert blocks and "just a paragraph" in blocks[0][0]
    assert blocks[0][2] == ""


def test_a_chunk_never_spans_two_sections_it_does_not_declare():
    blocks = [("first section body", None, "A > one"), ("second section body", None, "A > two")]
    packed = chunk_blocks(blocks, 900, 140)
    for chunk in packed:
        for word, section in (("first", "one"), ("second", "two")):
            if word in chunk["content"]:
                assert section in chunk["section_path"], \
                    f"chunk carries {word!r} text without declaring section {section!r}"


def test_packing_stops_at_the_chunk_size():
    blocks = [(f"section {index} " + "x" * 200, None, f"A > s{index}") for index in range(6)]
    packed = chunk_blocks(blocks, 500, 50)
    assert len(packed) > 1
    assert all(len(chunk["content"]) <= 500 for chunk in packed)


def test_a_section_longer_than_the_chunk_size_splits_and_keeps_its_path():
    blocks = [("word " * 500, None, "A > long")]
    packed = chunk_blocks(blocks, 400, 40)
    assert len(packed) > 1
    assert all(chunk["section_path"] == "A > long" for chunk in packed)


def test_a_packed_chunk_declares_every_section_it_covers():
    blocks = [("alpha", None, "Root > A"), ("beta", None, "Root > B"), ("gamma", None, "Root > C")]
    chunk = chunk_blocks(blocks, 900, 100)[0]
    assert chunk["section_path"] == "Root > A | B | C"
    assert chunk["sections"] == ["Root > A", "Root > B", "Root > C"]
    assert build_toc([chunk["sections"]]) == [
        {"path": "Root > A", "title": "A", "depth": 2},
        {"path": "Root > B", "title": "B", "depth": 2},
        {"path": "Root > C", "title": "C", "depth": 2},
    ]


def test_the_covered_sections_are_carried_as_data_not_parsed_from_the_label():
    """The label is ambiguous once a tail contains the separator itself.

    "A > B | C > d" could be two sections or three, so nothing may recover the
    list from the string; the list travels with the chunk.
    """
    blocks = [("one", None, "A > B"), ("two", None, "A > C > d")]
    chunk = chunk_blocks(blocks, 900, 100)[0]
    assert chunk["sections"] == ["A > B", "A > C > d"]
    assert [entry["path"] for entry in build_toc([chunk["sections"]])] == ["A > B", "A > C > d"]


def test_a_document_of_several_chapters_yields_a_complete_table_of_contents():
    """Each section appears once, whichever chunk ended up carrying it."""
    text = ("# Guide\n\n## One\n\nfirst body\n\n### Deep\n\ndeeper body\n\n"
            "## Two\n\nsecond body\n")
    chunks = chunk_blocks(extract("guide.md", text.encode("utf-8")), 900, 100)
    paths = [entry["path"] for entry in build_toc([chunk["sections"] for chunk in chunks])]
    assert paths == ["Guide", "Guide > One", "Guide > One > Deep", "Guide > Two"]


def test_front_matter_does_not_become_a_heading():
    text = "---\ntitle: HR policy\nauthor: Ops\n---\n\n# Real heading\n\nbody\n"
    assert [h.title for h in detect_markdown(text)] == ["Real heading"]


def test_setext_headings_still_work_outside_front_matter():
    text = "Big title\n=========\n\nbody\n\nSub\n---\n\nmore\n"
    assert [(h.level, h.title) for h in detect_markdown(text)] == [(1, "Big title"), (2, "Sub")]


def test_merging_shares_ancestry_once_and_keeps_a_single_path_intact():
    assert merge_section_paths(["A > B > x", "A > B > y"]) == "A > B > x | y"
    assert merge_section_paths(["A > B", "A > B"]) == "A > B"
    assert merge_section_paths(["A > x", "C > y"]) == "A > x | C > y"
    assert merge_section_paths(["", ""]) == ""


def test_a_structureless_document_chunks_exactly_as_before():
    text = "sentence. " * 300
    blocks = [(text, 1, None)]
    packed = chunk_blocks(blocks, 300, 50)
    assert len(packed) > 2
    assert all(chunk["page"] == 1 for chunk in packed)
    assert all(chunk["section_path"] == "" for chunk in packed)


def test_csv_and_json_carry_no_section_and_do_not_raise():
    for name, payload in (("t.csv", b"a,b\n1,2\n"), ("t.json", b'{"k": 1}')):
        chunks = chunk_blocks(extract(name, payload), 900, 100)
        assert chunks
        assert all(chunk["section_path"] == "" for chunk in chunks)


def test_markdown_ingest_produces_a_path_for_every_chunk():
    payload = ("# Guide\n\n## Leave\n\n### Sick\n\nup to eight days\n\n"
               "### Unpaid\n\nup to three months\n").encode("utf-8")
    chunks = chunk_blocks(extract("g.md", payload), 900, 100)
    assert chunks
    assert all(chunk["section_path"] for chunk in chunks)
    assert any("Sick" in chunk["section_path"] for chunk in chunks)


def test_pdf_uses_its_own_outline_in_preference_to_patterns():
    from eval.make_fixtures import build_pdf
    from pypdf import PdfReader, PdfWriter

    from app.structure import pdf_outline

    source = PdfReader(io.BytesIO(build_pdf([["Body of page one"], ["Body of page two"]])))
    writer = PdfWriter()
    for page in source.pages:
        writer.add_page(page)
    parent = writer.add_outline_item("Opening", 0)
    writer.add_outline_item("Nested detail", 1, parent=parent)
    buffer = io.BytesIO()
    writer.write(buffer)

    per_page = pdf_outline(PdfReader(io.BytesIO(buffer.getvalue())))
    assert per_page[1] == ["Opening"]
    assert per_page[2] == ["Opening", "Nested detail"]

    chunks = chunk_blocks(extract("outlined.pdf", buffer.getvalue()), 900, 100)
    assert any("Nested detail" in chunk["section_path"] for chunk in chunks)


def test_pdf_without_an_outline_falls_back_to_heading_patterns():
    from eval.make_fixtures import build_pdf

    payload = build_pdf([["Section 1. Network summary", "Shipments rose eleven percent."]])
    chunks = chunk_blocks(extract("plain.pdf", payload), 900, 100)
    assert any("Section 1" in chunk["section_path"] for chunk in chunks)


def test_a_malformed_outline_degrades_to_no_outline_rather_than_raising():
    from app.structure import pdf_outline

    class Broken:
        @property
        def outline(self):
            raise ValueError("corrupt outline")

    assert pdf_outline(Broken()) == {}


def test_detect_picks_the_right_strategy_per_suffix():
    markdown = "# Title\n\nbody\n"
    assert [h.title for h in detect(markdown, ".md")] == ["Title"]
    assert detect(markdown, ".csv") == []
    # A Markdown file numbering its sections in prose still gets a path.
    assert [h.title for h in detect("ماده ۱ — هدف\nمتن\n", ".md")] == ["ماده ۱ — هدف"]


def test_the_path_separator_is_not_produced_by_a_single_heading():
    text = "# Only\n\nbody\n"
    spans = section_paths(text, detect_markdown(text))
    assert PATH_SEPARATOR not in spans[-1][2]


@pytest.mark.parametrize("line", [
    "تبصره:", "ماده ۱۲:", "فصل اول:", "Article 5:", "Section 4.",
])
def test_a_heading_may_end_in_a_colon_or_a_full_stop(line):
    """Both are ordinary heading punctuation, and _clean_title strips them."""
    assert detect_plain(line + "\nbody\n"), f"{line!r} should still read as a heading"


def test_a_leading_horizontal_rule_is_not_front_matter():
    """Skipping to the next "---" would swallow every heading between them."""
    text = "---\n\n# Introduction\n\nbody\n\n## Setup\n\nmore\n\n---\n\n# Appendix\n\ntail\n"
    assert [h.title for h in detect_markdown(text)] == ["Introduction", "Setup", "Appendix"]


def test_front_matter_still_recognised_when_it_carries_keys():
    text = "---\ntitle: HR policy\nauthor: Ops\n---\n\n# Real heading\n\nbody\n"
    assert [h.title for h in detect_markdown(text)] == ["Real heading"]


def test_a_dotted_number_splits_sections_at_ingest_but_not_when_quoting():
    """"4.2.1 Rollback" and "7.4.0 fixes the leak" are indistinguishable.

    At ingest a false positive only adds a section boundary. When splitting a
    chunk to quote from it, a false positive makes the line unquotable, so the
    reader is never shown the sentence that answers.
    """
    from app.structure import looks_like_heading

    line = "7.4.0 fixes the memory leak"
    assert looks_like_heading(line) is True
    assert looks_like_heading(line, unambiguous=True) is False
    # An unmistakable heading is one under either reading.
    for certain in ("## Setup", "ماده ۳ — معامله کوچک"):
        assert looks_like_heading(certain, unambiguous=True) is True


def test_a_word_heading_joins_its_first_paragraph_rather_than_standing_alone():
    """A heading followed by an oversized paragraph must not chunk by itself."""
    import io as byte_io

    from docx import Document

    document = Document()
    document.add_heading("Rollback procedure", level=1)
    document.add_paragraph("word " * 400)
    buffer = byte_io.BytesIO()
    document.save(buffer)

    chunks = chunk_blocks(extract("runbook.docx", buffer.getvalue()), 300, 40)
    assert "Rollback procedure" in chunks[0]["content"]
    assert chunks[0]["content"].strip() != "Rollback procedure", \
        "a chunk holding only a heading carries no answer and still competes"


def test_back_to_back_word_headings_do_not_each_become_a_chunk():
    import io as byte_io

    from docx import Document

    document = Document()
    document.add_heading("Alpha", level=1)
    document.add_heading("Beta", level=2)
    document.add_paragraph("Body under beta.")
    buffer = byte_io.BytesIO()
    document.save(buffer)

    chunks = chunk_blocks(extract("nested.docx", buffer.getvalue()), 900, 100)
    assert len(chunks) == 1
    assert chunks[0]["section_path"] == "Alpha > Beta"
    assert "Body under beta." in chunks[0]["content"]


@pytest.mark.parametrize("line", [
    "Section 3 covers rollback and recovery.",
    "بند ۲ حقوق پایه افزایش می‌یابد.",
    "ماده ۵ حقوق را تعیین می‌کند.",
])
def test_a_sentence_opening_with_a_heading_keyword_is_still_a_sentence(line):
    """What follows the numbering decides, not how the line ends.

    "Section 4." is a heading and "Section 3 covers rollback." is not, and both
    end in a full stop; the first has nothing after its number, the second
    continues into a clause.
    """
    assert not detect_plain(line + "\nbody\n"), f"{line!r} is prose"


@pytest.mark.parametrize("line", [
    "تبصره ۱", "Chapter 4", "ماده ۱۲:", "Section 4.", "Section 1. Network summary",
    "فصل اول — کلیات", "ماده ۱ طرفین قرارداد",
])
def test_a_numbering_followed_by_nothing_a_separator_or_a_short_label_is_a_heading(line):
    assert detect_plain(line + "\nbody\n"), f"{line!r} should read as a heading"


def test_front_matter_keys_do_not_have_to_be_english():
    """A Persian-first corpus writes Persian keys, and its closing "---" would
    otherwise be read as a setext underline."""
    text = "---\nعنوان: راهنما\nنویسنده: واحد اداری\n---\n\n# سرفصل واقعی\n\nمتن\n"
    assert [h.title for h in detect_markdown(text)] == ["سرفصل واقعی"]


def test_consecutive_word_headings_never_strand_one_as_its_own_chunk():
    """A title above a chapter above a long paragraph is the failing shape."""
    import io as byte_io

    from docx import Document

    document = Document()
    document.add_heading("HR Policy", level=0)
    document.add_heading("Leave", level=1)
    document.add_paragraph("word " * 200)
    buffer = byte_io.BytesIO()
    document.save(buffer)

    chunks = chunk_blocks(extract("policy.docx", buffer.getvalue()), 300, 40)
    assert "HR Policy" in chunks[0]["content"]
    assert "word" in chunks[0]["content"], "no chunk may hold headings and nothing else"


@pytest.mark.parametrize("line", [
    "Chapter 3 Network security and incident response",
    "فصل سوم شرایط عمومی استخدام کارکنان دولت",
])
def test_an_unpunctuated_heading_may_still_name_its_subject(line):
    assert detect_plain(line + "\nbody\n"), f"{line!r} should read as a heading"


@pytest.mark.parametrize("line", [
    "ماده ۱۲ تعیین حقوق و مزایای کارکنان رسمی",
    "Chapter 7 Terms governing the supply of professional services",
])
def test_a_long_heading_without_punctuation_is_deliberately_given_up(line):
    """No length separates these from prose that opens with the same keyword.

    "Chapter 4 was written by the finance team" has a six-word tail too, and
    reading it as a heading fabricates a section path — which this module
    treats as worse than having none. So the cut falls on the conservative
    side, and a document that punctuates its headings keeps them at any length.
    """
    assert not detect_plain(line + "\nbody\n")
    keyword, number, *rest = line.split()
    for separator in ("—", ":"):
        punctuated = f"{keyword} {number} {separator} {' '.join(rest)}"
        assert detect_plain(punctuated + "\nbody\n"), f"{punctuated!r} should be a heading"


@pytest.mark.parametrize("line", [
    "Section 3 covers rollback and recovery.",
    "ماده ۵ حقوق را تعیین می‌کند.",
    "بند ۲ حقوق پایه افزایش می‌یابد.",
])
def test_both_heading_entry_points_agree_that_a_clause_is_prose(line):
    """looks_like_heading restated the rule instead of deferring to it, so a
    line detect_plain rejected was still split out as a heading when quoting."""
    from app.structure import looks_like_heading

    assert not detect_plain(line + "\n")
    assert looks_like_heading(line) is False
    assert looks_like_heading(line, unambiguous=True) is False


def test_sibling_word_headings_each_keep_their_own_path():
    """Merging them into the following section filed the first under the
    second's name and lost it from the table of contents."""
    import io as byte_io

    from docx import Document

    document = Document()
    document.add_heading("Alpha", level=1)
    document.add_heading("Beta", level=1)
    document.add_paragraph("Body under beta.")
    buffer = byte_io.BytesIO()
    document.save(buffer)

    chunk = chunk_blocks(extract("siblings.docx", buffer.getvalue()), 900, 100)[0]
    assert chunk["sections"] == ["Alpha", "Beta"]


def test_a_short_block_before_an_oversized_one_is_carried_into_it():
    """Stranding is a size question, so it is handled once for every format
    rather than per reader: a lone heading is a chunk with no answer in it."""
    blocks = [("Chapter one", None, "Chapter one"), ("word " * 200, None, "Chapter one > Detail")]
    chunks = chunk_blocks(blocks, 300, 40)
    assert "Chapter one" in chunks[0]["content"]
    assert "word" in chunks[0]["content"], "the heading must not be a chunk of its own"
    assert chunks[0]["sections"] == ["Chapter one", "Chapter one > Detail"]
    # Later windows are continuations of the long section alone.
    assert chunks[1]["sections"] == ["Chapter one > Detail"]


def test_sections_of_reads_the_list_and_falls_back_to_the_path():
    import json as json_module

    from app.structure import sections_of

    assert sections_of(json_module.dumps({"sections": ["A > B", "A > C"]}), "A > B | C") == ["A > B", "A > C"]
    assert sections_of("{}", "A > B") == ["A > B"]
    assert sections_of("not json", "A > B") == ["A > B"]
    assert sections_of(None, None) == []


def test_a_heading_never_ends_a_chunk_it_does_not_belong_to():
    """A chunk closing on a heading leaves that section's text in the next
    chunk without the words a reader would search for."""
    blocks = [("Alpha", None, "Alpha"), ("Alpha body. " * 18, None, "Alpha"),
              ("Sick leave", None, "Sick leave"),
              ("Up to eight days per year with a certificate.", None, "Sick leave")]
    chunks = chunk_blocks(blocks, 300, 40)
    for chunk in chunks:
        assert not chunk["content"].strip().endswith("Sick leave"), \
            "the heading must open the next chunk, not close this one"
    holder = next(c for c in chunks if "eight days" in c["content"])
    assert "Sick leave" in holder["content"]


def test_a_document_ending_on_a_heading_does_not_leave_it_alone():
    blocks = [("Alpha", None, "Alpha"), ("word " * 80, None, "Alpha"),
              ("Beta heading", None, "Beta heading")]
    chunks = chunk_blocks(blocks, 300, 40)
    assert chunks[-1]["content"].strip() != "Beta heading"
    assert "Beta heading" in chunks[-1]["content"]
    assert "Beta heading" in chunks[-1]["sections"]


def test_carried_text_keeps_its_own_page_and_declares_only_what_it_holds():
    """A window made entirely of carried text must not claim the long section,
    and must not be stamped with that section's page."""
    chunks = chunk_blocks([("X " * 145, 1, "Intro"), ("L " * 400, 2, "Long")], 300, 40)
    assert chunks[0]["page"] == 1
    assert chunks[0]["sections"] == ["Intro"]
    assert all(chunk["page"] == 2 for chunk in chunks[1:])
    assert all(chunk["sections"] == ["Long"] for chunk in chunks[1:])


def test_a_window_is_cited_by_the_page_supplying_most_of_it():
    """A heading carried across a page boundary must not drag the citation
    back with it: the reader is looking at the body, which is on page two."""
    chunks = chunk_blocks([("Chapter one", 1, "Chapter one"),
                           ("word " * 200, 2, "Chapter one > Detail")], 300, 40)
    assert chunks[0]["sections"] == ["Chapter one", "Chapter one > Detail"]
    assert all(chunk["page"] == 2 for chunk in chunks), "the body is on page two"

    # When the carried text is most of the window, its page is the right one.
    mostly_carried = chunk_blocks([("X " * 145, 1, "Intro"), ("L " * 400, 2, "Long")], 300, 40)
    assert mostly_carried[0]["page"] == 1
    assert mostly_carried[0]["sections"] == ["Intro"]


@pytest.mark.parametrize("line", [
    "Chapter 7 — Terms governing the supply of professional services",
    "فصل سوم — شرایط عمومی استخدام کارکنان دولت و نهادهای وابسته",
    "ماده ۱۲: تعیین حقوق و مزایای کارکنان رسمی و پیمانی",
])
def test_a_punctuated_heading_may_be_as_long_as_it_needs(line):
    """The word count guards the bare dotted-number rule, which has no
    punctuation to read; a heading that separates its title does."""
    assert detect_plain(line + "\nbody\n"), f"{line!r} should read as a heading"


def test_the_word_count_still_guards_the_dotted_number_rule():
    assert not detect_plain("3.5 percent of the fleet was replaced during the year\n")
    assert detect_plain("4.2.1 Rollback procedure\n")


@pytest.mark.parametrize("line", [
    "ماده ۵: حقوق و مزایای کارکنان رسمی توسط هیئت مدیره تعیین می شود.",
    "Section 4. This section describes the rollback and recovery procedure for the fleet.",
    "Chapter 4 - was written by the finance team in the last quarter of the year",
])
def test_a_separator_does_not_turn_a_sentence_into_a_heading(line):
    """A separator says a title follows, not that whatever follows is one."""
    assert not detect_plain(line + "\nbody\n"), f"{line!r} is prose"


def test_packing_respects_the_size_bound_even_with_no_body_anywhere():
    """A run of headings has nothing to attach to, so the bound is all there is.

    The end-of-document merge used to append them to the last chunk without
    checking, pushing it past chunk_size.
    """
    blocks = [(f"Heading number {index}", None, f"Heading number {index}") for index in range(29)]
    chunks = chunk_blocks(blocks, 200, 20)
    assert len(chunks) > 1
    assert all(len(chunk["content"]) <= 200 for chunk in chunks)


def test_headings_too_long_to_carry_still_open_the_section_they_precede():
    """Emitting them separately produced the body-less chunk this forbids."""
    headings = [(f"Chapter {index} — " + "title words here " * 4, None, f"Chapter {index}")
                for index in range(4)]
    chunks = chunk_blocks(headings + [("word " * 200, None, "Long")], 300, 40)
    for chunk in chunks:
        assert chunk["sections"], "no chunk may be filed under nothing"
    holder = next(c for c in chunks if "word word" in c["content"])
    assert "Long" in holder["sections"]
